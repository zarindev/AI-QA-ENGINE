"""Excel exports (openpyxl): the test-case workbook (Phase 3) and the full QA workbook (qa_workbook.py)."""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.storage.repository import Repository
from app.storage.schemas import Project, RequirementsDoc, Run, TestSuite

TESTCASES_XLSX = "exports/test_cases.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="0A0F1C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="E2E8F0")
BORDER = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
PRIORITY_FILL = {
    "P1": PatternFill("solid", fgColor="FEE2E2"),
    "P2": PatternFill("solid", fgColor="FEF3C7"),
    "P3": PatternFill("solid", fgColor="DBEAFE"),
    "P4": PatternFill("solid", fgColor="F1F5F9"),
}
STATUS_FILL = {
    "approved": PatternFill("solid", fgColor="DCFCE7"),
    "confirmed": PatternFill("solid", fgColor="DCFCE7"),
    "edited": PatternFill("solid", fgColor="DCFCE7"),
    "skipped": PatternFill("solid", fgColor="F1F5F9"),
    "rejected": PatternFill("solid", fgColor="F1F5F9"),
    "draft": PatternFill("solid", fgColor="FEF3C7"),
    "proposed": PatternFill("solid", fgColor="FEF3C7"),
}


def write_table(
    ws: Worksheet, headers: Sequence[str], rows: Iterable[Sequence[Any]], widths: Sequence[int]
) -> None:
    ws.append(list(headers))
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(list(row))
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment, cell.border = WRAP, BORDER
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 22


def _colour(ws: Worksheet, column: int, fills: dict[str, PatternFill]) -> None:
    for row in ws.iter_rows(min_row=2, min_col=column, max_col=column):
        fill = fills.get(str(row[0].value))
        if fill:
            row[0].fill = fill


def _cases_sheet(wb: Workbook, suite: TestSuite, title: str = "Test Cases", first: bool = False) -> Worksheet:
    ws = wb.active if first else wb.create_sheet(title)
    assert ws is not None
    ws.title = title
    rows = []
    for c in suite.cases:
        steps = "\n".join(
            f"{s.order}. {s.action}"
            + (f"  [{s.data}]" if s.data else "")
            + (f"\n   → {s.expected}" if s.expected else "")
            for s in c.steps
        )
        links = "; ".join(f"{k}: {', '.join(v)}" for k, v in c.links.items() if k != "pages")
        rows.append(
            [
                c.id,
                c.title,
                c.module,
                c.technique.replace("_", " "),
                c.type,
                c.priority,
                c.role,
                "\n".join(c.preconditions),
                steps,
                c.expected_result,
                links,
                ", ".join(c.links.get("pages", [])),
                "Full" if c.requires_full_mode else "Safe",
                c.status,
                c.source,
            ]
        )
    write_table(
        ws,
        [
            "ID",
            "Title",
            "Module",
            "Technique",
            "Type",
            "Priority",
            "Role",
            "Preconditions",
            "Steps",
            "Expected result",
            "Covers",
            "Pages",
            "Mode",
            "Status",
            "Source",
        ],
        rows,
        [14, 42, 16, 15, 12, 9, 13, 30, 70, 40, 26, 30, 8, 11, 12],
    )
    _colour(ws, 6, PRIORITY_FILL)
    _colour(ws, 14, STATUS_FILL)
    return ws


def _stories_sheet(wb: Workbook, req: RequirementsDoc) -> None:
    st = wb.create_sheet("User Stories")
    write_table(
        st,
        ["ID", "Role", "Module", "Feature", "Story", "Acceptance criteria", "Status"],
        [
            [
                s.id,
                s.role,
                s.module,
                s.feature,
                s.story,
                "\n".join(f"• {a}" for a in s.acceptance_criteria),
                s.status,
            ]
            for s in req.stories
        ],
        [9, 14, 16, 22, 55, 60, 11],
    )
    _colour(st, 7, STATUS_FILL)


def _rules_sheet(wb: Workbook, req: RequirementsDoc, title: str = "Business Rules") -> None:
    br = wb.create_sheet(title)
    write_table(
        br,
        ["ID", "Rule", "Condition", "Category", "Entity", "Source", "Confidence", "Status"],
        [
            [
                r.id,
                r.statement,
                r.condition,
                r.category,
                r.entity,
                r.source,
                round(r.confidence, 2),
                r.status,
            ]
            for r in req.rules
        ],
        [9, 50, 45, 14, 14, 30, 11, 11],
    )
    _colour(br, 8, STATUS_FILL)


def testcases_workbook(project: Project, run: Run, suite: TestSuite, req: RequirementsDoc | None) -> bytes:
    wb = Workbook()
    _cases_sheet(wb, suite, first=True)
    if req is not None:
        _stories_sheet(wb, req)
        _rules_sheet(wb, req)
    about = wb.create_sheet("About")
    for row in (
        ["Project", project.name],
        ["URL", project.url],
        ["Run", run.id],
        ["Cases", len(suite.cases)],
        ["Approved", sum(1 for c in suite.cases if c.status == "approved")],
        ["Generated by", "QA Pilot — test data is fake; records created by tests are prefixed QAP_"],
    ):
        about.append(row)
    about.column_dimensions["A"].width, about.column_dimensions["B"].width = 16, 70
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_testcases(
    repo: Repository, project: Project, run: Run, suite: TestSuite, req: RequirementsDoc | None
) -> Path:
    path = repo.run_path(run, TESTCASES_XLSX)
    repo.write_bytes(path, testcases_workbook(project, run, suite, req))
    return path
