"""The full QA workbook: Summary, Site Profile, Requirements, User Stories, Test Cases, Execution Results, Bugs,
Permission Matrix, Traceability — styled, colour-coded, frozen headers, filters, hyperlinks to evidence.

Evidence links are relative to the workbook (`exports/qa_report.xlsx` → `../artifacts/...`), so they open when the
run folder is kept together (or zipped).
"""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.data import ReportData
from app.reports.excel import (
    HEADER_FILL,
    HEADER_FONT,
    PRIORITY_FILL,
    _cases_sheet,
    _colour,
    _rules_sheet,
    _stories_sheet,
    write_table,
)
from app.storage.repository import Repository
from app.storage.schemas import ExecutionsDoc

QA_XLSX = "exports/qa_report.xlsx"

RESULT_FILL = {
    "pass": PatternFill("solid", fgColor="DCFCE7"),
    "fail": PatternFill("solid", fgColor="FEE2E2"),
    "blocked": PatternFill("solid", fgColor="FEF3C7"),
    "error": PatternFill("solid", fgColor="F1F5F9"),
    "not run": PatternFill("solid", fgColor="F1F5F9"),
}
SEVERITY_FILL = {
    "critical": PatternFill("solid", fgColor="FCA5A5"),
    "major": PatternFill("solid", fgColor="FED7AA"),
    "minor": PatternFill("solid", fgColor="FEF3C7"),
    "trivial": PatternFill("solid", fgColor="F1F5F9"),
}
MATRIX_FILL = {
    "allow": PatternFill("solid", fgColor="DCFCE7"),
    "deny": PatternFill("solid", fgColor="F1F5F9"),
    "untested": PatternFill("solid", fgColor="FFFFFF"),
    "HOLE": PatternFill("solid", fgColor="FCA5A5"),
}
LINK_FONT = Font(color="2563EB", underline="single")


def _link(ws: Worksheet, row: int, col: int, target: str, label: str) -> None:
    cell = ws.cell(row=row, column=col)
    cell.value = label
    cell.hyperlink = target
    cell.font = LINK_FONT


def _summary(wb: Workbook, d: ReportData) -> None:
    ws = wb.active
    assert ws is not None
    ws.title = "Summary"
    accent = d.branding.accent_color.lstrip("#").upper()
    ws["A1"] = f"{d.project.name} — QA report"
    ws["A1"].font = Font(size=18, bold=True, color=accent)
    ws["A2"] = (
        f"{d.project.url} · run {d.run.id} · {d.run.mode} mode · generated {d.generated:%Y-%m-%d %H:%M} UTC"
    )
    ws["A2"].font = Font(color="64748B")
    if d.branding.company_name:
        ws["A3"] = f"Prepared by {d.branding.company_name}" + (
            f" for {d.branding.prepared_for}" if d.branding.prepared_for else ""
        )
    counts, sev = d.counts(), d.severity_counts()
    q = d.quality
    rows: list[tuple[str, object]] = [
        (
            "Quality score",
            f"{q.score:g} / 100 (grade {q.grade})" if q and q.score is not None else "not measured",
        ),
        ("Domain", f"{d.profile.domain} · {d.profile.sub_type}" if d.profile else "—"),
        ("Pages explored", len(d.crawl.pages) if d.crawl else 0),
        ("Test cases run", len(d.result_list)),
        ("Passed", counts["pass"]),
        ("Failed", counts["fail"]),
        ("Blocked (Safe Mode / not reachable)", counts["blocked"]),
        ("Errors", counts["error"]),
        ("Open bugs", len(d.open_bugs)),
        ("  critical", sev["critical"]),
        ("  major", sev["major"]),
        ("  minor", sev["minor"]),
        ("  trivial", sev["trivial"]),
        ("Permission holes", d.matrix.holes if d.matrix else 0),
        ("Claude cost (tests)", f"${sum(r.cost_usd for r in d.result_list):.2f}"),
        ("Privacy blur", "on" if d.privacy else "off"),
    ]
    start = 5
    for i, (k, v) in enumerate(rows):
        ws.cell(row=start + i, column=1, value=k).font = Font(bold=not k.startswith("  "))
        ws.cell(row=start + i, column=2, value=v)
    if q:
        r = start + len(rows) + 1
        for j, h in enumerate(["Sub-score", "Score", "Passed / measured", "Open bugs", "Note"], 1):
            c = ws.cell(row=r, column=j, value=h)
            c.fill, c.font = HEADER_FILL, HEADER_FONT
        for i, s in enumerate(q.sub_scores, 1):
            ws.cell(row=r + i, column=1, value=s.label)
            ws.cell(row=r + i, column=2, value=s.score if s.score is not None else "—")
            ws.cell(row=r + i, column=3, value=f"{s.passed} / {s.executed}" if s.executed else "—")
            ws.cell(row=r + i, column=4, value=s.bugs)
            ws.cell(row=r + i, column=5, value=s.note)
        ws.cell(row=r + len(q.sub_scores) + 1, column=1, value=f"Formula: {q.formula}").font = Font(
            italic=True, color="64748B"
        )
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 36, 34
    ws.column_dimensions["C"].width, ws.column_dimensions["D"].width = 18, 11
    ws.column_dimensions["E"].width = 40


def _profile(wb: Workbook, d: ReportData) -> None:
    if not d.profile:
        return
    p = d.profile
    ws = wb.create_sheet("Site Profile")
    ws.append(["Domain", f"{p.domain} · {p.sub_type}"])
    ws.append(
        [
            "Confidence",
            f"{p.confidence:.0%} ({'Claude analysis' if p.method == 'ai' else 'offline heuristic'})",
        ]
    )
    ws.append(["Summary", p.summary])
    ws.append(["Roles", ", ".join(p.roles)])
    ws.append(["Modules", ", ".join(p.modules)])
    ws.append([])
    ws.append(["Feature", "Module", "Description", "Pages"])
    hdr = ws.max_row
    for cell in ws[hdr]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for f in p.features:
        ws.append([f.name, f.module, f.description, ", ".join(f.pages)])
    ws.append([])
    ws.append(["Entity", "Fields", "Operations", "Pages"])
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for e in p.entities:
        ws.append([e.name, ", ".join(f.name for f in e.fields), ", ".join(e.operations), ", ".join(e.pages)])
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[0].value in ("Domain", "Confidence", "Summary", "Roles", "Modules"):
            row[0].font = Font(bold=True)
    for col, w in zip("ABCD", (26, 30, 70, 40), strict=True):
        ws.column_dimensions[col].width = w


def _requirements(wb: Workbook, d: ReportData) -> None:
    if not d.req:
        return
    _rules_sheet(wb, d.req, title="Requirements")
    ws = wb["Requirements"]
    if d.req.workflows:
        ws.append([])
        ws.append(["Workflow", "Name", "States", "Forbidden transitions", "Roles"])
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        for w in d.req.workflows:
            forbidden = "; ".join(f"{t.from_state} → {t.to_state}" for t in w.transitions if not t.allowed)
            ws.append([w.id, w.name, " → ".join(w.states), forbidden or "—", ", ".join(w.roles)])
    _stories_sheet(wb, d.req)


def _results(wb: Workbook, d: ReportData, repo: Repository) -> None:
    ws = wb.create_sheet("Execution Results")
    rows = [
        [
            r.test_case_id,
            r.title,
            r.module,
            r.technique.replace("_", " "),
            r.priority,
            r.role,
            r.viewport,
            r.result,
            r.reproducibility,
            "yes" if r.flaky else "",
            r.method,
            r.reason,
            round(r.duration_ms / 1000, 1),
            round(r.cost_usd, 3),
            ", ".join(r.bug_ids),
            "",
        ]
        for r in d.result_list
    ]
    write_table(
        ws,
        [
            "Test case",
            "Title",
            "Module",
            "Technique",
            "Priority",
            "Role",
            "Viewport",
            "Result",
            "Repro",
            "Flaky",
            "Ran by",
            "Reason",
            "Seconds",
            "Cost $",
            "Bugs",
            "Video",
        ],
        rows,
        [16, 40, 14, 14, 8, 12, 10, 9, 7, 6, 12, 60, 8, 8, 14, 10],
    )
    _colour(ws, 5, PRIORITY_FILL)
    _colour(ws, 8, RESULT_FILL)
    for i, r in enumerate(d.result_list, 2):
        if repo.has_run_doc(d.run, f"executions/{r.test_case_id}.json"):
            attempts = repo.load_run_doc(d.run, f"executions/{r.test_case_id}.json", ExecutionsDoc).attempts
            if attempts and attempts[0].video:
                _link(ws, i, 16, f"../{attempts[0].video}", "video")


def _bugs(wb: Workbook, d: ReportData) -> None:
    ws = wb.create_sheet("Bugs")
    rows = [
        [
            b.id,
            b.title,
            b.module,
            b.severity,
            b.priority,
            b.status.replace("_", " "),
            round(b.confidence, 2),
            b.reproducibility,
            b.summary,
            "\n".join(b.preconditions),
            "\n".join(f"{i}. {s}" for i, s in enumerate(b.steps_to_reproduce, 1)),
            b.expected,
            b.actual,
            f"{b.environment.browser} · {b.environment.os} · {b.environment.viewport} · {b.environment.role}",
            b.environment.url,
            ", ".join(b.test_case_ids),
            ", ".join(b.links.get("rules", []) + b.links.get("stories", [])),
            "",
            "",
        ]
        for b in d.bugs
    ]
    write_table(
        ws,
        [
            "Bug",
            "Title",
            "Module",
            "Severity",
            "Priority",
            "Status",
            "Confidence",
            "Repro",
            "Summary",
            "Preconditions",
            "Steps to reproduce",
            "Expected",
            "Actual",
            "Environment",
            "URL",
            "Test cases",
            "Requirements",
            "Screenshot",
            "Clip",
        ],
        rows,
        [9, 44, 14, 9, 8, 12, 10, 7, 50, 30, 60, 40, 40, 36, 36, 18, 18, 12, 8],
    )
    _colour(ws, 4, SEVERITY_FILL)
    _colour(ws, 5, PRIORITY_FILL)
    for i, b in enumerate(d.bugs, 2):
        shot = b.annotated_screenshot or b.screenshot
        if shot:
            _link(ws, i, 18, f"../{shot}", "screenshot")
        if b.clip:
            _link(ws, i, 19, f"../{b.clip}", "clip")


def _matrix(wb: Workbook, d: ReportData) -> None:
    if not d.matrix or not d.matrix.roles:
        return
    m = d.matrix
    ws = wb.create_sheet("Permission Matrix")
    ws.append(["Screen", *m.roles])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    cells = {(c.role, c.page): c for c in m.cells}
    for page in m.pages:
        row = [page]
        for role in m.roles:
            c = cells.get((role, page))
            row.append("HOLE" if c and c.hole else (c.observed if c else ""))
        ws.append(row)
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            fill = MATRIX_FILL.get(str(cell.value))
            if fill:
                cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 44
    for i in range(2, len(m.roles) + 2):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 15
    ws.freeze_panes = "B2"
    ws.append([])
    ws.append(
        ["Legend: allow = can open · deny = refused · HOLE = opens but the role's menus never lead there"]
    )


def _traceability(wb: Workbook, d: ReportData) -> None:
    ws = wb.create_sheet("Traceability")
    rows = [
        [t.requirement, t.requirement_text, t.story, t.test_case, t.test_title, t.result, ", ".join(t.bugs)]
        for t in d.traceability()
    ]
    write_table(
        ws,
        ["Requirement", "Statement", "Stories", "Test case", "Test", "Result", "Bugs"],
        rows,
        [12, 50, 18, 18, 44, 9, 16],
    )
    _colour(ws, 6, RESULT_FILL)


def qa_workbook(repo: Repository, d: ReportData) -> bytes:
    wb = Workbook()
    _summary(wb, d)
    _profile(wb, d)
    _requirements(wb, d)
    if d.suite:
        _cases_sheet(wb, d.suite)
    _results(wb, d, repo)
    _bugs(wb, d)
    _matrix(wb, d)
    _traceability(wb, d)
    # sheet order required by the spec: Summary, Site Profile, Requirements, User Stories, Test Cases, ...
    order = [
        "Summary",
        "Site Profile",
        "Requirements",
        "User Stories",
        "Test Cases",
        "Execution Results",
        "Bugs",
        "Permission Matrix",
        "Traceability",
    ]
    wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export(repo: Repository, d: ReportData) -> Path:
    path = repo.run_path(d.run, QA_XLSX)
    repo.write_bytes(path, qa_workbook(repo, d))
    return path
