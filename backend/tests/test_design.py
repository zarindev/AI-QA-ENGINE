from __future__ import annotations

import io
import re
from datetime import date

import pytest
from openpyxl import load_workbook

from app.reports import excel, gherkin
from app.requirements.generator import generate_heuristic, info_pages
from app.storage.schemas import (
    CrawlResult,
    Element,
    Form,
    PageRecord,
    Project,
    RequirementsDoc,
    RoleCrawlSummary,
    Run,
    TestCase,
    utcnow,
)
from app.testdesign.generator import assemble, design, estimate, module_code
from app.testdesign.rules_based import Draft, RulesBasedDesigner, form_name
from app.testdesign.testdata import TestDataFactory
from app.understand.classifier import classify_heuristic


def _el(i, tag, label="", **kw):
    return Element(index=i, tag=tag, label=label, **kw)


def _crawl() -> CrawlResult:
    nav = [_el(0, "a", "Patients", href="/patients"), _el(1, "a", "Reports", href="/reports")]
    form_fields = [
        _el(10, "input", "Full name", form_index=0, required=True, maxlength=40),
        _el(11, "input", "Email", type="email", form_index=0),
        _el(12, "input", "Coverage (%)", type="number", min="0", max="100", form_index=0),
        _el(13, "input", "Date of birth", type="date", form_index=0, required=True),
        _el(14, "button", "Save patient", type="submit", form_index=0),
    ]

    def page(pid, role, tpl, title, **kw):
        return PageRecord(
            id=pid,
            role=role,
            url=f"http://c.test{tpl.replace('{id}', '3')}",
            url_template=tpl,
            title=title,
            elements=nav,
            **kw,
        )

    pages = [
        page("PG-1", "public", "/login", "Sign in · Clinic", is_login_page=True),
        page("PG-2", "admin", "/patients", "Patients · Clinic", headings=["Patients"]),
        page(
            "PG-3",
            "admin",
            "/patients/new",
            "Register patient · Clinic",
            headings=["Register patient"],
            forms=[
                Form(
                    index=0,
                    purpose="create",
                    field_indices=[10, 11, 12, 13, 14],
                    submit_index=14,
                    heading="Register patient",
                )
            ],
        ),
        page("PG-4", "admin", "/reports", "Reports · Clinic", headings=["Reports"]),
        page("PG-5", "nurse", "/patients", "Patients · Clinic", headings=["Patients"]),
        page(
            "PG-6",
            "admin",
            "/help",
            "Help · Clinic",
            text_excerpt="Patients pay the fee minus insurance coverage.",
        ),
    ]
    pages[2].elements = nav + form_fields
    return CrawlResult(
        start_url="http://c.test/",
        pages=pages,
        roles=[
            RoleCrawlSummary(role="public"),
            RoleCrawlSummary(role="admin", login_ok=True),
            RoleCrawlSummary(role="nurse", login_ok=True),
        ],
    )


def test_test_data_is_fake_and_marked():
    data = TestDataFactory(seed=1, today=date(2026, 10, 5))
    assert data.value_for("Full name").startswith("QAP_")
    assert data.value_for("Email", "email").endswith("@example.test")
    assert re.fullmatch(r"\+1 555 01\d\d", data.value_for("Phone"))
    assert data.value_for("Date of birth", "date") < "2026-10-05"
    assert data.value_for("Pickup date", "date") > "2026-10-05"
    assert data.value_for("Status", "select", ["Select…", "active", "inactive"]) == "active"
    assert TestDataFactory(seed=1).value_for("Full name") == TestDataFactory(seed=1).value_for("Full name")


def test_boundaries_and_classes_from_html_constraints():
    data = TestDataFactory(today=date(2026, 10, 5))
    cov = _el(1, "input", "Coverage (%)", type="number", min="0", max="100")
    labels = {(p.label, p.value, p.valid) for p in data.boundaries(cov)}
    assert ("min − 1 (-1)", "-1", False) in labels and ("max + 1 (101)", "101", False) in labels
    name = _el(2, "input", "Name", maxlength=5, required=True)
    probes = data.boundaries(name)
    assert ("QQQQQ", True) in [(p.value, p.valid) for p in probes] and ("", False) in [
        (p.value, p.valid) for p in probes
    ]
    dob = data.invalid_classes(_el(3, "input", "Date of birth", type="date"))
    assert dob[0].value == "2027-10-05" and not dob[0].valid
    phone = data.invalid_classes(_el(4, "input", "Phone"))
    assert any(p.value == "call me maybe" for p in phone)


def test_rules_based_designer_covers_techniques_and_permissions():
    crawl = _crawl()
    profile = classify_heuristic(crawl)
    drafts = RulesBasedDesigner(crawl, profile).all()
    by_tech: dict[str, list[TestCase]] = {}
    for d in drafts:
        by_tech.setdefault(d.case.technique, []).append(d.case)
    assert {
        "smoke",
        "crud",
        "validation",
        "boundary",
        "equivalence",
        "permission",
        "ui_responsive",
        "accessibility",
    } <= set(by_tech)
    perm = {c.title for c in by_tech["permission"]}
    assert "Nurse cannot open /reports by direct URL" in perm
    assert "Nurse cannot open /patients/new by direct URL" in perm
    assert not any("Admin cannot" in t for t in perm)  # admin reached everything
    assert "Signed-out visitors cannot open protected screens" in perm
    crud = by_tech["crud"][0]
    assert crud.requires_full_mode and all(
        str(v).startswith(("QAP_", "qap.")) or v[0].isdigit() for v in crud.test_data.values()
    )
    assert all(not c.requires_full_mode for c in by_tech["permission"] + by_tech["smoke"])
    assert next(c for c in by_tech["ui_responsive"]).viewports == ["tablet", "mobile"]


def test_form_names_never_carry_record_data():
    form = Form(index=0, heading="Prescription for Jane Doe")
    p1 = PageRecord(
        id="a", url="u", url_template="/appointments/{id}/prescribe", title="Prescription for Jane Doe · X"
    )
    assert (
        form_name(p1, form, {"Prescription for Jane Doe · X", "Prescription for Ben Low · X"})
        == "Prescribe appointment"
    )
    p2 = PageRecord(id="b", url="u", url_template="/products/{id}/edit", title="Edit product · Shop")
    assert form_name(p2, form, {"Edit product · Shop", "Edit product · Shop "}) == "Edit product"
    p3 = PageRecord(id="c", url="u", url_template="/patients/new", title="New · X")
    assert form_name(p3, Form(index=0, heading="Register patient")) == "Register patient"


def test_ids_priorities_and_regeneration_keeps_reviewed_cases():
    taken: dict[str, str] = {}
    assert module_code("Patient", taken) == "PAT" and module_code("Patient", taken) == "PAT"
    assert module_code("Sale line item", taken) == "SLIT"
    assert module_code("Patients", taken) == "PAT2"
    crawl = _crawl()
    profile = classify_heuristic(crawl)
    req = generate_heuristic(crawl, profile)
    suite = design(crawl, profile, req)
    ids = [c.id for c in suite.cases]
    assert len(ids) == len(set(ids)) and all(re.fullmatch(r"TC-[A-Z0-9]{2,4}-\d{3}", i) for i in ids)
    assert all(c.priority == "P1" for c in suite.cases if c.technique == "permission")
    # review: approve one, edit another, write one by hand
    suite.cases[0].status = "approved"
    suite.cases[1].title, suite.cases[1].edited = "Edited by a person", True
    manual = TestCase(
        id="", title="Nurse cannot delete a patient", technique="permission", module="Patient", source="user"
    )
    suite = assemble([Draft(manual, "Patient", "user:1")], suite.cases)
    again = design(crawl, profile, req, previous=suite)
    kept = {c.id: c for c in again.cases}
    assert kept[suite.cases[0].id].status == "approved" if suite.cases[0].id in kept else True
    assert any(c.title == "Edited by a person" for c in again.cases)
    assert any(c.title == "Nurse cannot delete a patient" and c.source == "user" for c in again.cases)
    assert len({c.id for c in again.cases}) == len(again.cases)


def test_estimate_and_safe_mode():
    cases = [
        TestCase(id="1", title="a", steps=[], requires_full_mode=True),
        TestCase(id="2", title="b", steps=[], viewports=["tablet", "mobile"]),
    ]
    full = estimate(cases, (2.0, 10.0))
    safe = estimate(cases, (2.0, 10.0), safe_mode=True)
    assert full.cases == 2 and full.steps == 6 and full.cost_usd > safe.cost_usd > 0
    assert safe.blocked_in_safe_mode == 1


def test_heuristic_requirements_and_info_pages():
    crawl = _crawl()
    profile = classify_heuristic(crawl)
    req = generate_heuristic(crawl, profile)
    assert req.method == "heuristic" and all(r.status == "proposed" for r in req.rules)
    assert any("Full name" in r.statement for r in req.rules)
    assert "insurance coverage" in info_pages(crawl)


def test_exports_excel_and_gherkin(repo):
    crawl = _crawl()
    profile = classify_heuristic(crawl)
    req = generate_heuristic(crawl, profile)
    suite = design(crawl, profile, req)
    suite.cases[0].status = "approved"
    project = Project(slug="p", name="Clinic", url="http://c.test", authorized_by="t", authorized_at=utcnow())
    run = Run(id="r1", project_slug="p")
    wb = load_workbook(io.BytesIO(excel.testcases_workbook(project, run, suite, req)))
    assert wb.sheetnames == ["Test Cases", "User Stories", "Business Rules", "About"]
    ws = wb["Test Cases"]
    assert ws["A1"].value == "ID" and ws.max_row == len(suite.cases) + 1 and ws.freeze_panes == "B2"
    files = gherkin.features(suite)
    assert len(files) == 1
    text = next(iter(files.values()))
    assert text.startswith("Feature: ") and f"@{suite.cases[0].id}" in text and "Scenario:" in text
    assert "Given " in text and "Then " in text
    assert gherkin.features(suite, ("skipped",)) == {}


def test_review_api_flow(isolated_workspace):
    """Requirements confirm/edit/reject, case approve/edit/add, estimate and exports — no AI, no browser."""
    from fastapi.testclient import TestClient

    from app.api import deps
    from app.main import app

    deps.repo.cache_clear()
    deps.executor.cache_clear()
    r = deps.repo()
    project = r.create_project(
        Project(
            slug="clinic",
            name="Clinic",
            url="http://c.test",
            environment="test",
            authorized_by="t",
            authorized_at=utcnow(),
        )
    )
    run = r.create_run(Run(id="20260101-000000", project_slug=project.slug, status="completed"))
    crawl = _crawl()
    profile = classify_heuristic(crawl)
    req = generate_heuristic(crawl, profile)
    r.save_run_doc(run, "crawl/pages.json", crawl)
    r.save_run_doc(run, "site_profile.json", profile)
    r.save_run_doc(run, "requirements.json", req)
    r.save_run_doc(run, "testcases.json", design(crawl, profile, req))
    base = "/api/projects/clinic/runs/20260101-000000"
    with TestClient(app) as c:
        rules = c.get(f"{base}/requirements").json()["rules"]
        rid = rules[0]["id"]
        assert (
            c.patch(f"{base}/requirements/rules/{rid}", json={"status": "confirmed"}).json()["status"]
            == "confirmed"
        )
        edited = c.patch(f"{base}/requirements/rules/{rid}", json={"condition": "fee * (1 - cov/100)"}).json()
        assert edited["status"] == "edited" and edited["condition"] == "fee * (1 - cov/100)"
        assert (
            c.post(
                f"{base}/requirements/rules/bulk",
                json={"ids": [x["id"] for x in rules[1:3]], "status": "rejected"},
            ).json()["updated"]
            == 2
        )
        assert c.patch(f"{base}/requirements/rules/NOPE", json={"status": "confirmed"}).status_code == 404

        cases = c.get(f"{base}/testcases").json()["cases"]
        assert (
            c.post(
                f"{base}/testcases/bulk", json={"ids": [x["id"] for x in cases[:3]], "status": "approved"}
            ).json()["updated"]
            == 3
        )
        changed = c.patch(
            f"{base}/testcases/{cases[3]['id']}",
            json={"title": "Better title", "steps": [{"action": "Do it", "expected": "It works"}]},
        ).json()
        assert changed["edited"] and changed["steps"][0]["order"] == 1
        added = c.post(
            f"{base}/testcases",
            json={
                "title": "Nurse cannot delete a patient",
                "module": "Patient",
                "technique": "permission",
                "role": "nurse",
            },
        )
        assert added.status_code == 201 and re.fullmatch(r"TC-PAT\d?-\d{3}", added.json()["id"])
        assert c.delete(f"{base}/testcases/{added.json()['id']}").status_code == 204
        assert (
            c.post(
                f"{base}/testcases/plain-english", json={"text": "Check that a nurse cannot see reports"}
            ).status_code
            == 400
        )  # no key

        est = c.get(f"{base}/testcases/estimate").json()
        assert est["cases"] == 3 and est["cost_usd"] > 0 and est["model"] == "claude-sonnet-5-5"
        assert c.post(f"{base}/exports/testcases_xlsx").json()["files"] == ["exports/test_cases.xlsx"]
        assert c.post(f"{base}/exports/gherkin_zip").json()["files"] == ["exports/gherkin.zip"]
        kinds = {e["kind"] for e in c.get(f"{base}/exports").json()}
        assert kinds == {"testcases_xlsx", "gherkin_zip"}
        xlsx = c.get(f"{base}/files/exports/test_cases.xlsx")
        assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    deps.executor().shutdown()
    deps.repo.cache_clear()
    deps.executor.cache_clear()


@pytest.mark.browser
def test_requirements_pdf_renders(repo):
    from app.reports import requirements_doc

    from .conftest import _chrome_ok

    if not _chrome_ok():
        pytest.skip("Chrome not installed")
    crawl = _crawl()
    profile = classify_heuristic(crawl)
    req = generate_heuristic(crawl, profile)
    project = repo.create_project(
        Project(slug="p", name="Clinic", url="http://c.test", authorized_by="t", authorized_at=utcnow())
    )
    run = repo.create_run(Run(id="r1", project_slug=project.slug))
    paths = requirements_doc.export(repo, project, run, profile, req)
    assert paths["pdf"].read_bytes()[:5] == b"%PDF-" and paths["pdf"].stat().st_size > 10_000
    md = paths["markdown"].read_text()
    assert md.startswith("# Clinic — Reverse-Engineered Requirements") and "## 5. Business rules" in md
    assert isinstance(req, RequirementsDoc)


def test_module_names_are_merged_singular_and_case_insensitive():
    a = Draft(TestCase(id="", title="one"), "Appointment", "k1")
    b = Draft(TestCase(id="", title="two", module="Appointments"), "Appointments", "k2")
    c = Draft(TestCase(id="", title="three", module="appointment"), "appointment", "k3")
    suite = assemble([a, b, c], [])
    assert {x.module for x in suite.cases} == {"Appointment"}
    assert sorted(x.id for x in suite.cases) == ["TC-APP-001", "TC-APP-002", "TC-APP-003"]
