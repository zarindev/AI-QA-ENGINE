from __future__ import annotations

import csv
import io
import py_compile
import zipfile

import pytest
from openpyxl import load_workbook

from app.execute.stage import BUGS_FILE, RESULTS_FILE
from app.reports import csv_export, pytest_export, qa_report, qa_workbook
from app.reports import data as report_data
from app.requirements.stage import REQUIREMENTS_FILE
from app.storage.schemas import (
    Bug,
    BugEnvironment,
    BugsDoc,
    BusinessRule,
    Execution,
    ExecutionsDoc,
    LocatorSet,
    Project,
    ReplayScript,
    ReplayStep,
    RequirementsDoc,
    ResultsDoc,
    Role,
    Run,
    StepResult,
    TestCase,
    TestRunResult,
    TestSuite,
    UserStory,
    utcnow,
)
from app.testdesign.stage import TESTCASES_FILE

from .conftest import requires_chrome


@pytest.fixture
def report(repo):
    project = repo.create_project(
        Project(
            slug="clinic",
            name="Clinic",
            url="http://clinic.test",
            authorized_by="pytest",
            authorized_at=utcnow(),
            environment="test",
            roles=[Role(name="admin", secret_ref="admin", login_url="http://clinic.test/login")],
        )
    )
    run = repo.create_run(Run(id="20260102-000000", project_slug="clinic", mode="full"))
    repo.save_run_doc(
        run,
        REQUIREMENTS_FILE,
        RequirementsDoc(
            rules=[
                BusinessRule(id="BR-001", statement="Patient pays fee × (1 − coverage)", status="confirmed")
            ],
            stories=[UserStory(id="US-001", role="admin", story="As an admin I bill visits")],
        ),
    )
    cases = [
        TestCase(
            id="TC-BIL-001",
            title="Bill an insured visit",
            module="Billing",
            technique="business_rule",
            role="admin",
            status="approved",
            requires_full_mode=True,
            expected_result='Patient pays "$60.00"',
            links={"rules": ["BR-001"], "stories": ["US-001"]},
        ),
        TestCase(
            id="TC-REP-001",
            title="Doctor cannot open /reports",
            module="Report",
            technique="permission",
            role="doctor",
            status="approved",
            steps=[{"order": 1, "action": "Open the URL", "data": "http://clinic.test/reports"}],
        ),
        TestCase(
            id="TC-PAT-002", title="Not run yet", module="Patients", technique="crud", status="approved"
        ),
    ]
    repo.save_run_doc(run, TESTCASES_FILE, TestSuite(cases=cases))
    repo.save_run_doc(
        run,
        RESULTS_FILE,
        ResultsDoc(
            results=[
                TestRunResult(
                    test_case_id="TC-BIL-001", technique="business_rule", result="fail", bug_ids=["BUG-001"]
                ),
                TestRunResult(test_case_id="TC-REP-001", technique="permission", result="pass"),
            ]
        ),
    )
    repo.save_run_doc(
        run,
        BUGS_FILE,
        BugsDoc(
            bugs=[
                Bug(
                    id="BUG-001",
                    title="Insurance coverage ignored",
                    severity="critical",
                    priority="P1",
                    module="Billing",
                    steps_to_reproduce=["Log in as admin", "Complete the visit"],
                    expected="Patient pays $60.00",
                    actual="Patient pays $120.00",
                    environment=BugEnvironment(browser="Chrome", viewport="1440×900 (desktop)", role="admin"),
                    test_case_ids=["TC-BIL-001"],
                    links={"rules": ["BR-001"]},
                ),
                Bug(id="BUG-002", title="Rejected one", status="rejected"),
            ]
        ),
    )
    ex = Execution(
        test_case_id="TC-BIL-001",
        result="fail",
        role="admin",
        expected="Patient pays $60.00 (50 % of $120.00)",
        actual="Patient pays $120.00",
        final_page_text="Invoice INV-7\nFee $120.00\nPatient pays $120.00",
        steps=[StepResult(order=1, action="navigate", input="http://clinic.test/visits/7")],
    )
    repo.save_run_doc(
        run, "executions/TC-BIL-001.json", ExecutionsDoc(test_case_id="TC-BIL-001", attempts=[ex])
    )
    repo.save_run_doc(
        run,
        "replay/TC-BIL-001.json",
        ReplayScript(
            test_case_id="TC-BIL-001",
            role="admin",
            steps=[
                ReplayStep(action="navigate", url="http://clinic.test/visits/7"),
                ReplayStep(
                    action="click",
                    target='[4] button "Complete visit"',
                    locators=LocatorSet(text="Complete visit", tag="button", id="complete"),
                ),
                ReplayStep(
                    action="type",
                    target='[5] input "Notes"',
                    input="QAP_note",
                    locators=LocatorSet(name="notes", tag="input"),
                ),
            ],
            checks=["Patient pays $60.00"],
        ),
    )
    return repo, report_data.load(repo, project, run)


def test_traceability_links_requirement_story_test_result_bug(report):
    _, d = report
    rows = {(t.requirement, t.test_case): t for t in d.traceability()}
    bil = rows[("BR-001", "TC-BIL-001")]
    assert bil.story == "US-001" and bil.result == "fail" and bil.bugs == ["BUG-001"]
    assert rows[("—", "TC-PAT-002")].result == "not run"


def test_workbook_has_all_sheets_in_order(report):
    repo, d = report
    wb = load_workbook(io.BytesIO(qa_workbook.qa_workbook(repo, d)))
    assert wb.sheetnames == [
        "Summary",
        "Requirements",
        "User Stories",
        "Test Cases",
        "Execution Results",
        "Bugs",
        "Traceability",
    ]  # Site Profile and Permission Matrix appear when exploration data exists
    bugs = wb["Bugs"]
    assert bugs["A2"].value == "BUG-001" and bugs.freeze_panes == "B2" and bugs.auto_filter.ref
    assert wb["Summary"]["A1"].value.startswith("Clinic")


def test_csv_exports_are_importable(report):
    _, d = report
    jira = list(csv.reader(io.StringIO(csv_export.jira_csv(d).decode("utf-8-sig"))))
    assert jira[0][:4] == ["Summary", "Issue Type", "Priority", "Description"]
    assert (
        len(jira) == 2 and jira[1][2] == "Highest" and "h3. Steps to reproduce" in jira[1][3]
    )  # rejected bug left out
    trello = list(csv.reader(io.StringIO(csv_export.trello_csv(d).decode("utf-8-sig"))))
    assert trello[1][0] == "[BUG-001] Insurance coverage ignored" and "critical" in trello[1][2]


def test_report_html_two_audiences_and_privacy_of_credentials(report):
    repo, d = report
    html = qa_report.render(repo, d, "full")
    assert "Executive summary" in html and "1 critical issue should be fixed before release" in html
    assert "Appendix B — Traceability" in html and "Insurance coverage ignored" in html
    assert "Rejected one" not in html
    single = qa_report.render(repo, d, "single", [d.bugs[0]])
    assert "Executive summary" not in single and "Patient pays $120.00" in single


def test_pytest_suite_compiles_and_asserts_expected_values(report, tmp_path):
    repo, d = report
    files = pytest_export.build(repo, d)
    assert {"conftest.py", "pages/base.py", "pages/billing.py", "tests/test_billing.py", "README.md"} <= set(
        files
    )
    for rel, content in files.items():
        if rel.endswith(".py"):
            p = tmp_path / rel.replace("/", "_")
            p.write_text(content if isinstance(content, str) else content.decode(), encoding="utf-8")
            py_compile.compile(str(p), doraise=True)
    billing = files["tests/test_billing.py"]
    assert "page.open('/visits/7')" in billing and "page.click(page.COMPLETE_VISIT_BUTTON)" in billing
    assert (
        "page.assert_contains('$60.00')" in billing
    )  # expected value missing on the page → test fails until fixed
    assert "@pytest.mark.full_mode" in billing
    report_tests = files["tests/test_report.py"]
    assert "page.access_refused('/reports')" in report_tests
    assert "pytest.skip('Not recorded yet" in files["tests/test_patients.py"]
    assert "Admin#" not in "".join(c for c in files.values() if isinstance(c, str))
    assert "QAP_ADMIN_PASSWORD=" in files[".env.example"]
    path = pytest_export.export(repo, d)
    assert "qa_pilot_suite/tests/test_billing.py" in zipfile.ZipFile(path).namelist()


def test_expectations_skip_negated_and_native_messages():
    case = TestCase(id="TC-1", title="t")
    ok = Execution(
        test_case_id="TC-1",
        result="pass",
        actual='Saw "Booking confirmed" and no "Hand over car" button; "Please fill out this field."',
        final_page_text="Booking confirmed\nHand over car",
    )
    assert pytest_export.expectations(case, ok)[0] == ["Booking confirmed"]


@requires_chrome
def test_single_bug_pdf(report):
    repo, d = report
    path = qa_report.export(repo, d, "single", "BUG-001")
    assert path.read_bytes()[:5] == b"%PDF-" and path.name == "BUG-001.pdf"
